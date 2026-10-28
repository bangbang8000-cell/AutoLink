# -*- coding: utf-8 -*-
"""B3（v5.5.6）物料清单 7 品类 + 机房几何 schema v2 专项测试

覆盖测试计划 T-B3-01 ~ T-B3-16 的后端可执行项（A16 / A17 / A18 / A21）：
  T-B3-01 7 品类一级节，顺序固定
  T-B3-02 节小计 = 该节明细行数量之和
  T-B3-03 全表总计 = 各节小计之和（公式引用数值小计行）
  T-B3-04 5 个 sheet 数量与名称不变
  T-B3-05 168 柜含 8 个 ac → 机柜 160 / 空调 8
  T-B3-06 柜型细分之和 === 机柜总数
  T-B3-07 未建模不给 0（写文本「未建模」）
  T-B3-08 机柜/空调走同一取值路径（后端同源断言：room_counts 与逐格遍历全等）
  T-B3-09 单排上限 24 硬校验
  T-B3-10 单排下限 6 与行数 1–26
  T-B3-11 rowCount × cabinetsPerRow === 面板总柜位
  T-B3-12 柜高 2.0 / 2.2 落盘
  T-B3-13 v1 项目可开，按默认值补齐并写回（迁移提示只一次）
  T-B3-14 通道宽度 1.2 / 1.35 / 1.6 落盘与读回一致
  T-B3-15 room_layout.v1.bak.json 内容 === 迁移前原 v1
  T-B3-16 v1 → v2 → 读取字段无损（往返门禁）
"""
import io
import json
from contextlib import redirect_stdout
from pathlib import Path

import pytest
from openpyxl import Workbook

import hwlist
import hwlist_sheets
from designer import NetworkDesignerV2
from room import (
    LAYOUT_FILENAME,
    PLACEHOLDER_AC,
    PLACEHOLDER_PILLAR,
    ROOM_SCHEMA_VERSION,
    ROOM_TYPE_EMPTY,
    V1_BACKUP_SUFFIX,
    RoomGeometry,
    RoomMatrix,
    load_room_layout,
    save_room_layout,
    validate_room_layout,
)

_TPL = (Path(__file__).resolve().parents[2] / 'template'
        / 'H100-128台-RoCE' / 'project_config.json')

ROW_8 = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']
COL_21 = list(range(1, 22))
# 160 个可建柜格的柜型分布（细分之和必须等于机柜总数）
TYPE_SPLIT = (('gpu', 64), ('compute', 48), ('storage', 32), ('network', 16))


@pytest.fixture(scope='module')
def design():
    buf = io.StringIO()
    with redirect_stdout(buf):
        return NetworkDesignerV2(str(_TPL))


def _matrix_168(ac=8):
    """8 行 × 21 列 = 168 柜位，其中 ac 个空调占位（T-B3-05 口径）"""
    m = RoomMatrix(ROW_8, COL_21, name='B3-TEST',
                   geometry=RoomGeometry(row_count=8, cabinets_per_row=21))
    positions = [f'{r}{c}' for r in ROW_8 for c in COL_21]
    i = 0
    for ctype, n in TYPE_SPLIT:
        for _ in range(n):
            m.set_type(positions[i], ctype)
            i += 1
    for pos in positions[i:i + ac]:
        m.set_placeholder(pos, PLACEHOLDER_AC)
    return m


def _lines_with_aux(design, matrix):
    """补齐「辅材」行：基础模板无授权/软件/服务行，7 品类全节需构造一行"""
    lines = hwlist.build_lines(design, matrix=matrix)
    aux = dict(lines[0])
    aux.update(device_id='misc_service_license', name='授权与软件（辅材）', model='',
               qty=1, unit='项', spare_ratio=0.0, spare_qty=0, total_qty=1,
               category='辅材', line_type='service', network_profile='机房配套',
               qty_basis='按项目授权项聚合')
    lines.append(aux)
    return lines


def _sheet(lines, variant='contract'):
    wb = Workbook()
    ws = wb.active
    hdr, subs = hwlist_sheets.write_hw_list(ws, lines, 'B3 用例', variant=variant)
    qcol = 7 if variant == 'contract' else 5
    return ws, hdr, subs, qcol


def _bands(ws):
    """按出现顺序返回 [(品类, band 行号)]"""
    out = []
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row, 1).value
        if isinstance(v, str) and '、' in v and v[0] in '一二三四五六七八九十':
            out.append((v.split('、', 1)[1], row))
    return out


def _detail_sum(ws, band_row, sub_row, qcol):
    """节内明细数量之和（跳过二级小计与「未建模」文本行）"""
    total = 0
    for row in range(band_row + 1, sub_row):
        if not isinstance(ws.cell(row, 1).value, int):
            continue
        v = ws.cell(row, qcol).value
        if isinstance(v, int):
            total += v
    return total


# ----------------------------------------------------------------------
# A16：一级 7 品类分节 + 双级小计与总计闭合
# ----------------------------------------------------------------------
class TestB3CategorySections:
    def test_t_b3_01_seven_category_bands_in_order(self, design):
        """T-B3-01：7 个一级节且顺序固定"""
        lines = _lines_with_aux(design, _matrix_168())
        ws, _, _, _ = _sheet(lines)
        cats = [c for c, _ in _bands(ws)]
        assert cats == list(hwlist.CATEGORY_ORDER), f'一级节顺序不符: {cats}'
        assert len(cats) == 7

    def test_t_b3_01_no_line_category_skipped(self, design):
        """T-B3-01 边界：无行的品类整节跳过，剩余节仍保持固定相对顺序"""
        lines = hwlist.build_lines(design, matrix=_matrix_168())
        ws, _, _, _ = _sheet(lines)
        cats = [c for c, _ in _bands(ws)]
        present = {ln['category'] for ln in lines}
        assert cats == [c for c in hwlist.CATEGORY_ORDER if c in present]
        assert '辅材' not in cats  # 基础模板无辅材行 → 整节跳过

    def test_t_b3_02_section_subtotal_equals_details(self, design):
        """T-B3-02：每节小计 === 该节明细行数量之和"""
        lines = _lines_with_aux(design, _matrix_168())
        ws, _, subs, qcol = _sheet(lines)
        band_rows = dict(_bands(ws))
        assert [c for c, _, _ in subs] == list(hwlist.CATEGORY_ORDER)
        for cat, sub_row, numeric in subs:
            assert numeric, f'{cat} 节应为数值小计'
            expected = sum(ln['qty'] for ln in lines
                           if ln['category'] == cat and not ln.get('not_modeled'))
            assert _detail_sum(ws, band_rows[cat], sub_row, qcol) == expected, \
                f'{cat} 节明细之和与预期不符'

    def test_t_b3_03_total_equals_sum_of_subtotals(self, design):
        """T-B3-03：总计公式 === 引用各节数值小计之和"""
        lines = _lines_with_aux(design, _matrix_168())
        ws, _, subs, qcol = _sheet(lines)
        numeric_rows = [r for _, r, ok in subs if ok]
        total_rows = [r for r in range(1, ws.max_row + 1)
                      if ws.cell(r, 3).value == '总计（数量清单 · 不含价格）']
        assert len(total_rows) == 1, '应有且只有一行表末总计'
        formula = ws.cell(total_rows[0], qcol).value
        assert formula == '=' + '+'.join(f'G{r}' for r in numeric_rows)
        # 数值闭合：总计应等于全部非未建模明细之和
        expected = sum(ln['qty'] for ln in lines if not ln.get('not_modeled'))
        assert sum(_detail_sum(ws, dict(_bands(ws))[c], r, qcol)
                   for c, r, ok in subs if ok) == expected

    def test_t_b3_04_sheet_structure_unchanged(self, design):
        """T-B3-04：5 个 sheet 数量与名称不变"""
        lines = hwlist.build_lines(design, matrix=_matrix_168())
        stat = hwlist.summarize(lines)
        wb = hwlist_sheets.build_workbook(lines, 'B3 用例', stat=stat)
        assert wb.sheetnames == ['设计汇总', '组网清单（简版）', '组网清单（契约版）',
                                 '组网示意图', '设计口径']

    def test_t_b3_02_simple_variant_subtotal(self, design):
        """T-B3-02：简版同样按节小计闭合（E 列为数量）"""
        lines = hwlist.build_lines(design, matrix=_matrix_168())
        ws, _, subs, qcol = _sheet(lines, variant='simple')
        band_rows = dict(_bands(ws))
        for cat, sub_row, ok in subs:
            if not ok:
                continue
            expected = sum(ln['qty'] for ln in lines if ln['category'] == cat)
            assert _detail_sum(ws, band_rows[cat], sub_row, qcol) == expected


# ----------------------------------------------------------------------
# A17：机柜/空调矩阵真值 + 未建模边界
# ----------------------------------------------------------------------
class TestB3RoomCounts:
    def test_t_b3_05_cabinets_160_ac_8(self, design):
        """T-B3-05：168 柜含 8 个 ac → 机柜 160 / 空调 8"""
        lines = hwlist.build_lines(design, matrix=_matrix_168())
        cab = [ln for ln in lines if ln['category'] == '机柜']
        ac = [ln for ln in lines if ln['category'] == '空调']
        assert sum(ln['qty'] for ln in cab) == 160
        assert len(ac) == 1 and ac[0]['qty'] == 8
        assert all(ln['name'].startswith('机柜 · ') for ln in cab)
        assert cab[0]['network_profile'] == hwlist.NET_ROOM[1]
        assert ac[0]['network_profile'] == hwlist.NET_ROOM[1]

    def test_t_b3_06_cabinet_type_split_closes(self, design):
        """T-B3-06：柜型细分之和 === 机柜总数"""
        m = _matrix_168()
        counts = hwlist.room_counts(m)
        assert sum(v for _, v in counts['by_type']) == counts['cabinets'] == 160
        by_type = dict(counts['by_type'])
        assert by_type == {'gpu': 64, 'compute': 48, 'storage': 32, 'network': 16}
        lines = hwlist.build_lines(design, matrix=m)
        assert sum(ln['qty'] for ln in lines if ln['category'] == '机柜') == counts['cabinets']

    def test_t_b3_05_pillar_not_counted(self):
        """T-B3-05 边界：柱子占位既不计机柜也不计空调；空白格不计"""
        m = RoomMatrix(['A', 'B'], list(range(1, 7)))
        m.set_type('A1', 'gpu')
        m.set_placeholder('A2', PLACEHOLDER_PILLAR)
        counts = hwlist.room_counts(m)
        assert counts['cabinets'] == 1
        assert counts['ac'] == 0
        assert counts['cells'] == 12

    def test_t_b3_07_not_modeled_is_text_not_zero(self, design):
        """T-B3-07：矩阵缺失/为空 → 「未建模」文本，不出现 0"""
        assert hwlist.room_counts(None) is None
        assert hwlist.room_counts({}) is None

        lines = hwlist.build_lines(design, matrix=None)
        room = [ln for ln in lines if ln['category'] in ('机柜', '空调')]
        assert len(room) == 1
        assert room[0]['not_modeled'] is True
        assert room[0]['qty'] == 0 and room[0]['unit'] == '未建模'

        ws, _, subs, qcol = _sheet(lines)
        cab_band = dict(_bands(ws))['机柜']
        cab_sub = [r for c, r, _ in subs if c == '机柜'][0]
        assert ws.cell(cab_sub, qcol).value == '未建模'
        for row in range(cab_band + 1, cab_sub + 1):
            v = ws.cell(row, qcol).value
            assert not (isinstance(v, int) and v == 0), '未建模节不得出现 0'
        total_row = [r for r in range(1, ws.max_row + 1)
                     if ws.cell(r, 3).value == '总计（数量清单 · 不含价格）'][0]
        numeric_rows = [r for _, r, ok in subs if ok]
        assert ws.cell(total_row, qcol).value == '=' + '+'.join(f'G{r}' for r in numeric_rows)

    def test_t_b3_07_unmodeled_warning(self, design):
        """T-B3-07：未建模进入口径告警"""
        warns = hwlist.warnings_of(hwlist.build_lines(design, matrix=None))
        assert any('机房矩阵未建模' in w for w in warns)

    def test_t_b3_08_single_source_path(self):
        """T-B3-08：room_counts 与逐格遍历全等；RoomMatrix 与落盘 dict 结果一致"""
        m = _matrix_168()
        counts = hwlist.room_counts(m)
        cabinets = ac = 0
        for cell in m.cells.values():
            if cell.placeholder == PLACEHOLDER_AC:
                ac += 1
            elif cell.placeholder is None and (
                    cell.cabinet_id is not None or cell.type != ROOM_TYPE_EMPTY):
                cabinets += 1
        assert (counts['cabinets'], counts['ac']) == (cabinets, ac) == (160, 8)
        # room_layout.json 的 dict 形态走同一取值逻辑
        from_dict = json.loads(json.dumps(m.to_dict(), ensure_ascii=False))
        assert hwlist.room_counts(from_dict) == counts

    def test_t_b3_08_seeded_cabinets_counted(self):
        """T-B3-08：已落柜（有 cabinetId、type 未标）同样计入机柜真值"""
        m = RoomMatrix(['A'], [1, 2, 3, 4, 5, 6])
        for pos, cid in (('A1', 101), ('A2', 102)):
            m.place_cabinet(pos, cid)
        counts = hwlist.room_counts(m)
        assert counts['cabinets'] == 2
        assert counts['by_type'] == [('empty', 2)]


# ----------------------------------------------------------------------
# A18：几何参数校验（行 1–26 / 每排 6–24 / 闭合 / 柜高 / 通道）
# ----------------------------------------------------------------------
class TestB3GeometryValidation:
    def test_t_b3_09_per_row_upper_bound_24(self):
        """T-B3-09：填 25 被拒并提示「单排上限 24 柜」"""
        errs = RoomGeometry(row_count=8, cabinets_per_row=25).validate(8 * 25)
        assert any('单排上限 24 柜' in e for e in errs)
        m = RoomMatrix(ROW_8, list(range(1, 26)))
        assert any('单排上限 24 柜' in e for e in m.validate())

    def test_t_b3_10_per_row_lower_bound_6(self):
        """T-B3-10：填 5 被拒并提示「单排下限 6 柜」"""
        errs = RoomGeometry(row_count=8, cabinets_per_row=5).validate(8 * 5)
        assert any('单排下限 6 柜' in e for e in errs)

    def test_t_b3_10_row_count_range(self):
        """T-B3-10：行数 0 与 27 被拒（合法区间 1–26）"""
        for rows in (0, 27):
            errs = RoomGeometry(row_count=rows, cabinets_per_row=6).validate(None)
            assert any('机柜行数需在 1–26 行' in e for e in errs), f'行数 {rows} 未被拒'
        for rows in (1, 26):
            assert RoomGeometry(row_count=rows, cabinets_per_row=6).validate(None) == []

    def test_t_b3_11_slots_must_match_panel(self):
        """T-B3-11：rowCount × cabinetsPerRow === 面板总柜位"""
        m = _matrix_168()
        assert m.geometry.row_count * m.geometry.cabinets_per_row == m.size == 168
        assert m.validate() == []
        m.geometry.cabinets_per_row = 20
        errs = m.validate()
        assert any('与面板总柜位 168 不一致' in e for e in errs)

    def test_t_b3_12_cabinet_height_persisted(self, tmp_path):
        """T-B3-12：42U / 47U 切换后 cabinetHeight 分别落盘 2.0 / 2.2"""
        path = tmp_path / LAYOUT_FILENAME
        for height in (2.0, 2.2):
            m = _matrix_168()
            m.geometry.cabinet_height = height
            save_room_layout(str(path), m)
            data = json.loads(path.read_text(encoding='utf-8'))
            assert data['schemaVersion'] == ROOM_SCHEMA_VERSION
            assert data['geometry']['cabinetHeight'] == pytest.approx(height)
            loaded = load_room_layout(str(path))
            assert loaded.geometry.cabinet_height == pytest.approx(height)

    def test_t_b3_12_default_geometry_from_panel(self):
        """T-B3-12 边界：未显式传几何 → 按面板行列补齐默认值"""
        m = RoomMatrix(ROW_8, COL_21)
        assert m.geometry.row_count == 8
        assert m.geometry.cabinets_per_row == 21
        assert m.geometry.cabinet_width == 0.6
        assert m.geometry.cabinet_depth == 1.2
        assert m.geometry.cabinet_height == 2.0
        assert m.validate() == []

    def test_t_b3_09_illegal_height_and_aisle(self):
        """T-B3-09 关联：柜高与通道仅允许枚举值"""
        errs = RoomGeometry(cabinet_height=2.1, aisle_width=1.5).validate(168)
        assert any('机柜高度仅支持 2.0 m（42U）/ 2.2 m（47U）' in e for e in errs)
        assert any('通道宽度仅支持 1.2 / 1.35 / 1.6 m' in e for e in errs)


# ----------------------------------------------------------------------
# A21：v1 兼容（备份 / 迁移 / 往返门禁）
# ----------------------------------------------------------------------
def _v1_payload():
    cells = []
    for r in ('A', 'B', 'C', 'D'):
        for c in range(1, 7):
            cells.append({
                'row': r,
                'col': c,
                'type': 'gpu' if c <= 2 else ROOM_TYPE_EMPTY,
                'placeholder': PLACEHOLDER_AC if (r == 'D' and c == 6) else None,
                'cabinetId': c if r == 'A' else None,
            })
    return {
        'schemaVersion': 1,
        'name': 'v1 机房',
        'rows': ['A', 'B', 'C', 'D'],
        'cols': [1, 2, 3, 4, 5, 6],
        'cells': cells,
    }


def _write_v1(path):
    payload = _v1_payload()
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    return payload


class TestB3SchemaCompatibility:
    def test_t_b3_13_v1_opens_and_migrates_once(self, tmp_path):
        """T-B3-13：v1 可开、按默认值补齐并写回；迁移提示只出现一次"""
        path = tmp_path / LAYOUT_FILENAME
        _write_v1(path)

        m = load_room_layout(str(path))
        assert m.migrated_from_v1 is True
        assert m.geometry.row_count == 4
        assert m.geometry.cabinets_per_row == 6
        assert m.geometry.cabinet_height == 2.0

        on_disk = json.loads(path.read_text(encoding='utf-8'))
        assert on_disk['schemaVersion'] == ROOM_SCHEMA_VERSION
        assert on_disk['geometry'] == m.geometry.to_dict()

        again = load_room_layout(str(path))
        assert again.migrated_from_v1 is False, 'v2 布局不应再次触发迁移提示'

    def test_t_b3_14_aisle_width_roundtrip(self, tmp_path):
        """T-B3-14：通道宽度 1.2 / 1.35 / 1.6 落盘值与读回值一致"""
        path = tmp_path / LAYOUT_FILENAME
        for width in (1.2, 1.35, 1.6):
            m = _matrix_168()
            m.geometry.aisle_width = width
            save_room_layout(str(path), m)
            data = json.loads(path.read_text(encoding='utf-8'))
            assert data['geometry']['aisleWidth'] == pytest.approx(width)
            assert load_room_layout(str(path)).geometry.aisle_width == pytest.approx(width)

    def test_t_b3_15_backup_equals_original_v1(self, tmp_path):
        """T-B3-15：room_layout.v1.bak.json 内容 === 迁移前原 v1"""
        path = tmp_path / LAYOUT_FILENAME
        original = _write_v1(path)
        original_bytes = path.read_bytes()

        load_room_layout(str(path))

        bak = tmp_path / (Path(LAYOUT_FILENAME).stem + V1_BACKUP_SUFFIX)
        assert bak.exists()
        assert bak.read_bytes() == original_bytes
        assert json.loads(bak.read_text(encoding='utf-8')) == original

        # 二次读取（已是 v2）不得覆盖备份
        bak_bytes = bak.read_bytes()
        load_room_layout(str(path))
        assert bak.read_bytes() == bak_bytes

    def test_t_b3_16_v1_to_v2_roundtrip_lossless(self, tmp_path):
        """T-B3-16：v1 → v2 → 读取字段无损 + .bak 存在（CI 门禁）"""
        path = tmp_path / LAYOUT_FILENAME
        original = _write_v1(path)

        first = load_room_layout(str(path)).to_dict()
        assert first['schemaVersion'] == ROOM_SCHEMA_VERSION
        assert first['name'] == original['name']
        assert first['rows'] == original['rows']
        assert first['cols'] == original['cols']

        def _key(c):
            return (c['row'], c['col'])

        assert {_key(c): (c['type'], c['placeholder'], c['cabinetId']) for c in first['cells']} == \
               {_key(c): (c['type'], c['placeholder'], c['cabinetId']) for c in original['cells']}

        second = load_room_layout(str(path)).to_dict()
        assert second == first, 'v2 往返后字段应完全无损'
        assert validate_room_layout(second) == []
        assert (tmp_path / (Path(LAYOUT_FILENAME).stem + V1_BACKUP_SUFFIX)).exists()

    def test_t_b3_16_v2_roundtrip_keeps_geometry(self, tmp_path):
        """T-B3-16：v2 自身往返不丢几何字段"""
        path = tmp_path / LAYOUT_FILENAME
        m = _matrix_168()
        m.geometry.cabinet_height = 2.2
        m.geometry.aisle_width = 1.6
        save_room_layout(str(path), m)
        loaded = load_room_layout(str(path))
        assert loaded.geometry.to_dict() == m.geometry.to_dict()
        assert loaded.migrated_from_v1 is False
